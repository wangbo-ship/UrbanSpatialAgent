# -*- coding: utf-8 -*-
"""生成「城市治理 · 空间问答与研判助手」设计文档（分文件 DOCX）"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent / "报告文档"
PROJECT = "城市治理 · 空间问答与研判助手"
SUBTITLE = "个人项目 · 设计文档"


def set_run_font(run, size=12, bold=False, color=None, font_name="宋体"):
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    run.font.name = font_name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), font_name)
    if font_name in ("黑体", "微软雅黑", "楷体"):
        rFonts.set(qn("w:ascii"), "Microsoft YaHei")
        rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    else:
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")


def init_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    return doc


def add_title(doc, text, level=0):
    if level == 0:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size=22, bold=True, font_name="黑体")
        p.paragraph_format.space_after = Pt(12)
        return p
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        set_run_font(run, size=16 if level == 1 else 14, bold=True, font_name="黑体")
    return heading


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=12, color=RGBColor(0x55, 0x55, 0x55), font_name="楷体")
    p.paragraph_format.space_after = Pt(18)
    return p


def add_para(doc, text, first_indent=True, bold=False):
    p = doc.add_paragraph()
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=bold, font_name="宋体")
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.left_indent = Cm(0.74 + level * 0.5)
    run = p.add_run(text)
    set_run_font(run, size=12, font_name="宋体")
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=12, font_name="宋体")
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=11, bold=True, font_name="黑体")
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10.5, font_name="宋体")
    doc.add_paragraph()
    return table


def add_meta_footer(doc, doc_name):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run(
        f"文档名称：{doc_name}\n项目名称：{PROJECT}\n文档类型：个人项目设计文档（分册）\n版本：V1.0"
    )
    set_run_font(run, size=10.5, color=RGBColor(0x66, 0x66, 0x66), font_name="宋体")


def cover_block(doc, doc_title):
    add_title(doc, PROJECT, 0)
    add_subtitle(doc, SUBTITLE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(doc_title)
    set_run_font(run, size=18, bold=True, font_name="黑体")
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(24)
    add_para(
        doc,
        "本文档为个人项目设计文档分册，与系统可运行演示、应用案例说明配套使用。"
        "内容围绕城市治理垂直场景，阐述基于大语言模型、AIGC 与智能体技术的空间问答与研判方案。",
        first_indent=True,
    )


# ---------------------------------------------------------------------------
# 1. 项目概述
# ---------------------------------------------------------------------------
def build_overview():
    doc = init_doc()
    cover_block(doc, "01 项目概述")

    add_title(doc, "一、项目背景", 1)
    add_para(
        doc,
        "随着新型智慧城市建设推进，城市治理业务对“空间感知—态势研判—辅助决策”的一体化能力提出更高要求。"
        "传统 GIS 系统功能强大，但操作门槛高、分析链路长，业务人员往往需要依赖专业分析师才能完成缓冲分析、"
        "叠加统计、设施可达性评估等常见空间研判任务。与此同时，大语言模型（LLM）、人工智能生成内容（AIGC）"
        "与智能体（Intelligent Agent）技术快速发展，使得“用自然语言驱动专业工具链”成为可能。",
    )
    add_para(
        doc,
        "本项目面向城市治理垂直场景，设计并开发「空间问答与研判助手」。系统以 Web 端为载体，融合 GIS 空间分析、"
        "知识库检索与 Agent 工具编排，实现“一句话提出研判问题—自动调用空间工具—地图可视化呈现—生成研判报告”的闭环。"
        "项目强调可演示、可落地、有明确行业价值，探索人工智能在城市治理空间研判中的落地路径。",
    )

    add_title(doc, "二、项目定位与目标", 1)
    add_para(doc, "项目定位：城市治理领域的 Spatial Agent（空间智能体）应用系统。", bold=True)
    add_para(doc, "总体目标如下：")
    add_numbered(doc, "降低空间分析门槛：业务人员用自然语言即可完成常见空间研判。")
    add_numbered(doc, "提升研判效率：将缓冲、叠加、统计、可达性等流程自动化编排。")
    add_numbered(doc, "结果可解释：展示 Agent 工具调用链路，保证分析过程透明可追溯。")
    add_numbered(doc, "支持案例演示：提供固定 Demo 剧本与样例数据，便于稳定复现与展示。")
    add_numbered(doc, "形成完整交付：设计文档 + 可运行软件 + 演示效果三者齐备。")

    add_title(doc, "三、应用场景边界", 1)
    add_para(doc, "本期聚焦以下城市治理子场景：")
    add_bullet(doc, "公共服务设施可达性与覆盖评估（学校、医院、公园、消防站点等）。")
    add_bullet(doc, "投诉热点/事件点周边环境研判（周边设施、人口/建筑密度代理指标）。")
    add_bullet(doc, "规划合规与安全距离校核（如拟建点与敏感目标缓冲区关系）。")
    add_bullet(doc, "自然语言空间问答与研判报告自动生成。")
    add_para(
        doc,
        "不在本期核心范围内的能力（可作为后续扩展说明）：高精度三维孪生、实时视频汇聚、"
        "全量城市级大数据中台建设、移动端外业采集专用 App。",
    )

    add_title(doc, "四、预期成果与交付物", 1)
    add_table(
        doc,
        ["交付物", "形式", "说明"],
        [
            ["设计文档", "DOCX 分册", "需求分析、设计思路、功能模块、工具平台等"],
            ["可运行系统", "Web 应用", "Vue 前端 + Python 后端，本地/局域网可演示"],
            ["样例数据", "GeoJSON/CSV", "脱敏城市设施与事件点样例"],
            ["演示剧本", "文档+预设问答", "2～3 个固定场景保证演示稳定"],
            ["技术说明", "README", "环境、启动、接口与依赖说明"],
        ],
    )

    add_title(doc, "五、创新点摘要", 1)
    add_bullet(doc, "“自然语言 → GIS 工具链”的 Agent 范式，而非简单的地图聊天框。")
    add_bullet(doc, "对话区与地图区联动，分析结果即时落图，演示直观。")
    add_bullet(doc, "工具调用轨迹可视化，满足“智能体”可解释性与过程可追溯的需求。")
    add_bullet(doc, "RAG 融合城市治理规范/手册，回答兼顾空间计算与政策知识。")
    add_bullet(doc, "技术栈贴近工程实践：Vue3 + FastAPI + LangChain + GeoPandas/PostGIS。")

    add_meta_footer(doc, "01 项目概述")
    path = OUT_DIR / "01_项目概述.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 2. 需求分析
# ---------------------------------------------------------------------------
def build_requirements():
    doc = init_doc()
    cover_block(doc, "02 需求分析")

    add_title(doc, "一、需求来源与调研方法", 1)
    add_para(
        doc,
        "需求来源于智慧城市场景中城管、规划、应急、公共服务等业务对空间研判的共性痛点。"
        "本项目通过文献与政策材料梳理、典型业务访谈假设、同类产品能力对比，并结合"
        "“垂直场景 + 大模型/AIGC/智能体 + 可运行系统”的落地目标，形成需求清单。",
    )
    add_para(doc, "调研与分析方法包括：")
    add_bullet(doc, "业务痛点归纳：操作门槛、多系统切换、结果解释成本高。")
    add_bullet(doc, "场景用例法：以用户故事描述“谁在什么情况下要完成什么目标”。")
    add_bullet(doc, "优先级矩阵：按业务价值、技术可行性、演示可视性排序。")
    add_bullet(doc, "非功能约束：演示时长可控、本机可运行、数据脱敏、稳定性。")

    add_title(doc, "二、用户角色分析", 1)
    add_table(
        doc,
        ["角色", "典型职责", "核心诉求"],
        [
            ["业务研判人员", "受理投诉、设施评估、巡查问题研判", "快速得到周边分析结论与地图"],
            ["规划/治理专员", "合规校核、覆盖盲区识别", "多条件空间叠加与报告输出"],
            ["值班指挥人员", "突发事件临机研判", "一句话问出影响范围与相关设施"],
            ["系统管理员", "数据导入、账号与配置", "样例数据管理、模型与密钥配置"],
            ["产品体验者", "了解与试用系统能力", "一键剧本、稳定输出、可解释链路"],
        ],
    )

    add_title(doc, "三、业务痛点", 1)
    add_numbered(doc, "专业 GIS 软件学习成本高，基层业务人员难以独立完成缓冲/叠加分析。")
    add_numbered(doc, "跨系统查询（地图、表格、文档）导致研判链路长、结果难汇总。")
    add_numbered(doc, "口头汇报依赖人工整理，缺乏标准化研判报告。")
    add_numbered(doc, "分析结果“黑盒”，领导追问“怎么算出来的”时难以快速解释。")
    add_numbered(doc, "规范条文与空间数据割裂，回答常只见地图不见依据。")

    add_title(doc, "四、功能性需求", 1)

    add_title(doc, "4.1 自然语言空间问答", 2)
    add_para(doc, "系统应支持用户以自然语言提出空间相关问题，例如：")
    add_bullet(doc, "“某某小区周边 500 米内有多少所学校和医院？”")
    add_bullet(doc, "“以投诉点为中心做 1 公里缓冲，统计落入缓冲区内的公园数量。”")
    add_bullet(doc, "“拟建点是否落入消防站 8 分钟可达覆盖范围？”")
    add_para(
        doc,
        "系统需理解意图、识别地点/半径/图层要素类型，并自动编排相应 GIS 工具完成计算，"
        "最后以自然语言答复并联动地图展示。",
    )

    add_title(doc, "4.2 空间分析能力", 2)
    add_table(
        doc,
        ["能力项", "需求描述", "优先级"],
        [
            ["缓冲分析", "按距离生成缓冲区并统计落入要素", "P0"],
            ["空间查询", "点/线/面相交、包含、邻近查询", "P0"],
            ["叠加统计", "多图层叠加后计数/分类汇总", "P0"],
            ["距离/可达性", "直线距离、简易服务范围评估", "P0"],
            ["热点聚合", "事件点聚类或核密度示意", "P1"],
            ["路径分析", "基于路网的最短路径（可扩展）", "P2"],
            ["栅格分析", "遥感/热力栅格解读（可扩展）", "P2"],
        ],
    )

    add_title(doc, "4.3 地图可视化与交互", 2)
    add_bullet(doc, "底图切换、缩放平移、图层开关。")
    add_bullet(doc, "分析结果自动叠加：缓冲区、命中要素高亮、标注统计结果。")
    add_bullet(doc, "地图点击反查：选中要素后可继续追问。")
    add_bullet(doc, "对话与地图双向联动：问答结果驱动视野定位。")

    add_title(doc, "4.4 智能体编排与可解释性", 2)
    add_bullet(doc, "Agent 自动选择并调用 GIS/检索工具。")
    add_bullet(doc, "展示工具调用序列（工具名、关键参数、耗时、结果摘要）。")
    add_bullet(doc, "支持失败重试与友好错误提示（如地名无法匹配）。")

    add_title(doc, "4.5 知识库问答（RAG）", 2)
    add_bullet(doc, "导入城市治理相关规范、操作手册、历史案例摘要。")
    add_bullet(doc, "回答中引用知识片段，并与空间结果结合生成综合研判。")

    add_title(doc, "4.6 研判报告生成（AIGC）", 2)
    add_bullet(doc, "一键生成结构化研判报告（背景、数据、方法、结论、建议）。")
    add_bullet(doc, "支持导出 Markdown/Word/PDF（至少实现一种）。")
    add_bullet(doc, "报告中嵌入关键统计表与地图截图说明（演示版可先文字+数据表）。")

    add_title(doc, "4.7 数据与系统管理", 2)
    add_bullet(doc, "样例图层导入（学校、医院、公园、投诉点等）。")
    add_bullet(doc, "地名/兴趣点别名库维护（提高问答命中率）。")
    add_bullet(doc, "模型接口、密钥、温度等参数配置。")
    add_bullet(doc, "演示剧本管理：预设问题与期望分析链路。")

    add_title(doc, "五、非功能性需求", 1)
    add_table(
        doc,
        ["类别", "要求"],
        [
            ["性能", "常规缓冲统计在样例数据规模下 3～10 秒内返回（含模型推理）"],
            ["可用性", "核心演示流程无需专业 GIS 操作知识"],
            ["可靠性", "提供离线样例与预设问答，降低演示时网络与模型波动影响"],
            ["安全性", "密钥不入库明文展示；演示数据脱敏；接口鉴权可简可扩"],
            ["可维护性", "前后端分离；工具注册化；配置外置"],
            ["可扩展性", "新增 GIS 工具无需改前端主流程"],
            ["兼容性", "主流 Chrome/Edge 浏览器；Windows 开发与演示环境"],
            ["可演示性", "单机或局域网可运行完整系统"],
        ],
    )

    add_title(doc, "六、用户故事（节选）", 1)
    add_para(
        doc,
        "US-01：作为城管研判人员，我希望输入“某投诉点 500 米内学校与医院数量”，"
        "以便快速判断周边公共服务配套情况。",
    )
    add_para(
        doc,
        "US-02：作为规划专员，我希望系统自动生成缓冲分析过程说明，"
        "以便向领导解释结论依据。",
    )
    add_para(
        doc,
        "US-03：作为值班人员，我希望用一句话查询拟建点与敏感目标距离是否达标，"
        "以便临机做出合规提示。",
    )
    add_para(
        doc,
        "US-04：作为产品体验者，我希望一键运行 Demo 剧本并看到地图与报告，"
        "以便快速了解系统完整能力。",
    )

    add_title(doc, "七、需求优先级与分期", 1)
    add_para(doc, "一期（核心必达）：", bold=True)
    add_bullet(doc, "自然语言问答 + 缓冲/查询/统计 + 地图联动 + 工具轨迹 + 报告草稿 + 样例数据。")
    add_para(doc, "二期（增强）：", bold=True)
    add_bullet(doc, "热点分析、多轮指代消解增强、报告导出 PDF、管理员后台完善。")
    add_para(doc, "三期（扩展）：", bold=True)
    add_bullet(doc, "路网可达、栅格/遥感解读、多 Agent 分工、移动端轻量入口。")

    add_title(doc, "八、约束与假设", 1)
    add_bullet(doc, "演示数据为脱敏样例，不代表真实行政区划精度。")
    add_bullet(doc, "一期可达性以直线距离/简单缓冲区近似，不承诺与真实路况一致。")
    add_bullet(doc, "大模型调用依赖可用 API；需提供模拟/降级应答策略。")
    add_bullet(doc, "坐标系默认 WGS84/Web Mercator，分析距离需投影到合适度量坐标系。")

    add_title(doc, "九、验收标准（需求级）", 1)
    add_numbered(doc, "至少完成 3 个演示用例端到端跑通（提问→分析→落图→答复/报告）。")
    add_numbered(doc, "Agent 工具调用轨迹对用户可见。")
    add_numbered(doc, "设计文档齐全，与实现模块对应。")
    add_numbered(doc, "本机可启动前端与后端，在浏览器完成演示。")
    add_numbered(doc, "对无法识别的地点给出明确提示，不出现静默失败。")

    add_meta_footer(doc, "02 需求分析")
    path = OUT_DIR / "02_需求分析.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 3. 设计思路
# ---------------------------------------------------------------------------
def build_design():
    doc = init_doc()
    cover_block(doc, "03 设计思路")

    add_title(doc, "一、总体设计理念", 1)
    add_para(
        doc,
        "本系统的核心设计理念是：把 GIS 专业能力封装为可被智能体调用的工具，"
        "把大语言模型作为“意图理解与流程编排大脑”，把 Web 地图作为“结果表达舞台”。"
        "用户不需要记忆缓冲区、叠加分析等菜单路径，只需用业务语言描述研判目标；"
        "系统负责把目标翻译为工具序列，并保证计算真实发生在 GIS 引擎中，而不是由模型“编造坐标与统计结果”。",
    )
    add_para(doc, "三条原则：", bold=True)
    add_bullet(doc, "计算可信：空间几何与统计由 GIS 库执行，LLM 不负责数值真相。")
    add_bullet(doc, "过程可见：工具调用链可展示，便于解释与复盘。")
    add_bullet(doc, "演示优先：固定剧本 + 样例数据，确保展示过程可控。")

    add_title(doc, "二、问题抽象：从业务问题到 Agent 任务", 1)
    add_para(
        doc,
        "城市治理空间研判问题可抽象为四元组：目标对象（地点/事件）、分析动作（缓冲/查询/统计）、"
        "约束条件（半径、图层类型、阈值）、输出形态（地图、数字、报告）。"
        "Agent 的职责是完成槽位填充、工具选择、参数校验、结果汇总与自然语言生成。",
    )
    add_table(
        doc,
        ["业务表述", "结构化意图", "工具序列示例"],
        [
            ["某小区 500m 内学校数", "POI定位→缓冲→图层筛选→计数", "geocode → buffer → spatial_filter → count"],
            ["投诉点周边设施清单", "事件定位→缓冲→多图层相交", "locate_event → buffer → intersect_layers"],
            ["拟建点是否过近敏感目标", "两点距离/缓冲相交判定", "project → distance/buffer → intersect → judge"],
        ],
    )

    add_title(doc, "三、系统分层架构思路", 1)
    add_para(doc, "采用前后端分离 + AI 编排层 + GIS 能力层：")
    add_numbered(doc, "表现层（Vue3）：对话、地图、轨迹面板、报告预览。")
    add_numbered(doc, "接入层（FastAPI）：鉴权、会话、流式输出、静态资源。")
    add_numbered(doc, "智能编排层（LangChain/LangGraph）：意图理解、工具选择、多步规划、记忆。")
    add_numbered(doc, "能力层：GIS Tools、RAG Retriever、Report Generator。")
    add_numbered(doc, "数据层：GeoJSON/PostGIS、向量库、配置与日志。")
    add_para(
        doc,
        "该分层保证“换模型不换工具、换底图不换 Agent、换场景主要改提示词与数据”，"
        "有利于后续继续扩展到农业、应急等其他垂直场景。",
    )

    add_title(doc, "四、关键技术路径", 1)

    add_title(doc, "4.1 为何选择 Web 而非 App", 2)
    add_para(
        doc,
        "本项目目标是在计算机上完整演示可运行系统。Web 端部署与展示成本最低，地图组件生态成熟，"
        "便于浏览器直接体验。App 更适合外业采集，非本期主路径；可在设计文档中保留扩展规划。",
    )

    add_title(doc, "4.2 LangChain 在本项目中的角色", 2)
    add_para(
        doc,
        "LangChain 用于组织 Prompt、工具（Tools）、检索器与 Agent 执行循环。"
        "本项目将缓冲、投影、相交、统计、检索等能力注册为 Tool；"
        "由 Agent 根据用户问题决定调用顺序。复杂多步流程可进一步用 LangGraph 固化状态机"
        "（规划→分析→校验→报告），降低自由 Agent 的不稳定。",
    )

    add_title(doc, "4.3 “防幻觉”设计", 2)
    add_bullet(doc, "禁止模型直接给出未经工具计算的面积/距离/数量结论。")
    add_bullet(doc, "工具返回结构化 JSON，再由模型润色措辞。")
    add_bullet(doc, "坐标分析前强制投影到度量坐标系，避免经纬度单位误用。")
    add_bullet(doc, "地名无法匹配时明确失败，引导用户点选地图或更换名称。")

    add_title(doc, "4.4 人机交互设计思路", 2)
    add_para(
        doc,
        "采用“左对话、右地图、底/侧轨迹”布局：用户提问后，先看到思考/工具步骤，再看到地图更新，"
        "最后看到总结与“生成报告”按钮。该顺序符合研判认知：先方法、再证据、后结论。",
    )

    add_title(doc, "五、数据与坐标系策略", 1)
    add_para(
        doc,
        "演示数据建议使用脱敏后的城市局部样例（学校、医院、公园、消防站、投诉事件点）。"
        "存储可采用 GeoJSON（轻量易交付）或 PostGIS（更贴近真实工程）。"
        "输入坐标统一为 EPSG:4326；涉及米制缓冲/面积时，动态选择合适的投影（如 UTM）或使用地理库的测地能力。",
    )

    add_title(doc, "六、安全与合规设计思路", 1)
    add_bullet(doc, "演示数据脱敏，避免真实隐私与精确敏感坐标。")
    add_bullet(doc, "API Key 仅存环境变量或本地配置，不进入前端。")
    add_bullet(doc, "对用户输入做长度与注入防护；工具参数白名单校验。")
    add_bullet(doc, "日志记录工具调用，不记录完整密钥。")

    add_title(doc, "七、演示设计思路", 1)
    add_para(doc, "建议准备三个固定剧本：")
    add_numbered(doc, "公共服务覆盖：小区周边学校/医院统计。")
    add_numbered(doc, "投诉研判：事件点缓冲内公园/设施清单。")
    add_numbered(doc, "合规距离：拟建点与敏感目标缓冲相交判定。")
    add_para(
        doc,
        "每个剧本提供一键按钮，减少操作误差；同时保留自由问答入口体现智能体通用性。",
    )

    add_title(doc, "八、设计权衡说明", 1)
    add_table(
        doc,
        ["议题", "选择", "理由"],
        [
            ["Agent 自由度", "工具约束 + 可选图工作流", "兼顾灵活与稳定"],
            ["地图引擎", "Leaflet 优先", "轻量、易集成、够演示"],
            ["空间库", "GeoPandas 起步，可接 PostGIS", "交付快且可升级"],
            ["报告", "先 Markdown/DOCX", "实现成本可控"],
            ["多模态", "一期不做强依赖", "聚焦空间 Agent 主线"],
        ],
    )

    add_title(doc, "九、小结", 1)
    add_para(
        doc,
        "设计思路可概括为：垂直场景牵引、Agent 编排核心、GIS 工具保真、地图表达落地、文档与演示闭环。"
        "该思路兼顾创新性与完整性，也为后续工程化留下清晰演进路径。",
    )

    add_meta_footer(doc, "03 设计思路")
    path = OUT_DIR / "03_设计思路.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 4. 功能模块
# ---------------------------------------------------------------------------
def build_modules():
    doc = init_doc()
    cover_block(doc, "04 功能模块")

    add_title(doc, "一、功能模块总览", 1)
    add_para(
        doc,
        "系统按“业务可见功能”与“支撑能力”划分模块。前端以场景化交互呈现，后端以服务与工具形式实现。"
        "模块之间通过 API 与事件（地图更新、轨迹推送、报告生成）解耦。",
    )
    add_table(
        doc,
        ["模块编号", "模块名称", "主要职责", "优先级"],
        [
            ["M1", "用户与会话模块", "登录（可选）、会话历史、上下文记忆", "P1"],
            ["M2", "智能对话模块", "自然语言输入、流式答复、多轮追问", "P0"],
            ["M3", "空间智能体编排模块", "意图理解、工具选择、多步执行", "P0"],
            ["M4", "GIS 分析工具模块", "缓冲、相交、统计、距离等", "P0"],
            ["M5", "地图可视化模块", "底图、图层、结果叠加、定位", "P0"],
            ["M6", "知识库 RAG 模块", "规范文档检索与引用", "P1"],
            ["M7", "研判报告模块", "结构化报告生成与导出", "P0"],
            ["M8", "数据管理模块", "样例图层、地名库、导入导出", "P0"],
            ["M9", "演示剧本模块", "一键 Demo、预设问答", "P0"],
            ["M10", "系统配置与监控模块", "模型配置、日志、健康检查", "P1"],
        ],
    )

    add_title(doc, "二、模块详细说明", 1)

    add_title(doc, "2.1 M2 智能对话模块", 2)
    add_para(doc, "功能点：")
    add_bullet(doc, "文本输入框、快捷问题、停止生成。")
    add_bullet(doc, "流式显示模型回复与中间状态。")
    add_bullet(doc, "消息气泡区分用户、助手、系统提示。")
    add_bullet(doc, "支持基于上一轮结果的追问（如“把半径改成 1 公里再算一次”）。")
    add_para(doc, "输入/输出：", bold=True)
    add_para(doc, "输入为用户文本与可选地图选点；输出为答复文本、结构化结果引用、地图指令。")

    add_title(doc, "2.2 M3 空间智能体编排模块", 2)
    add_para(
        doc,
        "本模块是系统“大脑”。基于 LangChain Agent（或 LangGraph 工作流）维护状态："
        "当前问题、已解析槽位、已调用工具、中间几何、最终结论。",
    )
    add_para(doc, "核心子功能：")
    add_bullet(doc, "意图分类：空间统计 / 合规判定 / 知识问答 / 报告生成 / 闲聊拒识。")
    add_bullet(doc, "槽位抽取：地点、半径、图层、阈值、输出类型。")
    add_bullet(doc, "工具规划与执行：按依赖顺序调用，失败则修正参数或追问用户。")
    add_bullet(doc, "轨迹记录：供前端展示与日志审计。")

    add_title(doc, "2.3 M4 GIS 分析工具模块", 2)
    add_table(
        doc,
        ["工具名", "功能", "主要参数", "返回"],
        [
            ["geocode_place", "地名解析为坐标/面", "name", "geometry/lonlat"],
            ["buffer_geometry", "生成缓冲区", "geometry, distance_m", "polygon"],
            ["intersect_layers", "图层相交筛选", "aoi, layer", "feature collection"],
            ["count_features", "要素计数", "features", "count"],
            ["summarize_by_type", "分类汇总", "features, field", "table"],
            ["distance_between", "距离计算", "g1, g2", "meters"],
            ["judge_within", "是否落入范围", "target, area", "bool + reason"],
        ],
    )
    add_para(
        doc,
        "实现上可用 GeoPandas/Shapely/PyProj；工具对外统一 JSON Schema，便于 Agent 调用与前端解析。",
    )

    add_title(doc, "2.4 M5 地图可视化模块", 2)
    add_bullet(doc, "加载底图与业务图层（学校、医院、公园、事件点等）。")
    add_bullet(doc, "接收后端 map_actions：fitBounds、addGeoJSON、highlight、setOpacity。")
    add_bullet(doc, "图例、图层管理、要素弹窗属性展示。")
    add_bullet(doc, "地图选点作为后续问答的空间上下文。")

    add_title(doc, "2.5 M6 知识库 RAG 模块", 2)
    add_para(
        doc,
        "对治理规范、设施配置标准、内部操作手册进行切分、向量化与检索。"
        "当问题同时涉及“怎么算”和“标准怎么规定”时，Agent 可并行调用 GIS 工具与检索器，"
        "再综合生成答复。",
    )

    add_title(doc, "2.6 M7 研判报告模块", 2)
    add_para(doc, "报告结构建议：")
    add_numbered(doc, "研判背景与问题描述")
    add_numbered(doc, "使用数据与范围")
    add_numbered(doc, "分析方法与工具链路")
    add_numbered(doc, "主要发现（表格+关键指标）")
    add_numbered(doc, "结论与建议")
    add_numbered(doc, "附录：工具调用明细")
    add_para(doc, "支持在对话结束后一键生成，并可再次润色。")

    add_title(doc, "2.7 M8 数据管理模块", 2)
    add_bullet(doc, "图层列表、字段预览、启用/禁用。")
    add_bullet(doc, "上传 GeoJSON/CSV（含经纬度）。")
    add_bullet(doc, "地名别名维护。")
    add_bullet(doc, "数据校验：坐标系、空几何、重复 ID。")

    add_title(doc, "2.8 M9 演示剧本模块", 2)
    add_para(
        doc,
        "内置演示按钮，点击后自动填充问题并执行完整链路。"
        "剧本配置可用 YAML/JSON 描述：问题文本、期望工具序列、成功判据。",
    )

    add_title(doc, "三、模块间协作流程（典型）", 1)
    add_para(doc, "以“某小区 500 米内学校数量”为例：")
    add_numbered(doc, "M2 接收用户问题并创建会话消息。")
    add_numbered(doc, "M3 解析意图与槽位，规划工具序列。")
    add_numbered(doc, "M4 执行 geocode → buffer → intersect → count。")
    add_numbered(doc, "M5 根据返回几何与要素集合更新地图。")
    add_numbered(doc, "M6（可选）补充设施配置相关规范条文。")
    add_numbered(doc, "M2 输出自然语言结论；用户可触发 M7 生成报告。")
    add_numbered(doc, "全程轨迹写入日志并由前端轨迹面板展示。")

    add_title(doc, "四、前端页面功能映射", 1)
    add_table(
        doc,
        ["页面/区域", "对应模块", "说明"],
        [
            ["对话区", "M2/M3", "提问、追问、流式答复"],
            ["地图区", "M5", "结果落图与交互选点"],
            ["轨迹抽屉", "M3", "工具调用可视化"],
            ["报告页/弹窗", "M7", "报告预览导出"],
            ["数据管理页", "M8", "样例与图层管理"],
            ["演示入口", "M9", "一键剧本"],
            ["设置页", "M10", "模型与系统参数"],
        ],
    )

    add_title(doc, "五、接口粒度建议（摘要）", 1)
    add_bullet(doc, "POST /api/chat ：对话与 Agent 执行（可 SSE）。")
    add_bullet(doc, "GET /api/layers ：图层列表。")
    add_bullet(doc, "POST /api/layers/upload ：数据上传。")
    add_bullet(doc, "POST /api/report ：报告生成。")
    add_bullet(doc, "GET /api/demos ：演示剧本列表。")
    add_bullet(doc, "GET /api/health ：健康检查。")

    add_title(doc, "六、功能验收清单", 1)
    add_bullet(doc, "P0 模块全部可在演示环境运行。")
    add_bullet(doc, "至少 3 条演示问答输出正确统计并落图。")
    add_bullet(doc, "工具轨迹至少展示工具名与关键参数。")
    add_bullet(doc, "报告可生成并包含结论段落。")
    add_bullet(doc, "异常地名有明确提示。")

    add_meta_footer(doc, "04 功能模块")
    path = OUT_DIR / "04_功能模块.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 5. 工具平台介绍
# ---------------------------------------------------------------------------
def build_platform():
    doc = init_doc()
    cover_block(doc, "05 工具平台介绍")

    add_title(doc, "一、平台定位", 1)
    add_para(
        doc,
        "「城市治理 · 空间问答与研判助手」采用“应用前端 + AI 编排中间层 + GIS/数据能力后端”的平台化组织方式。"
        "本节介绍开发与运行所依赖的主要工具、框架、平台组件及其在系统中的作用，"
        "便于理解技术选型合理性与可复现性。",
    )

    add_title(doc, "二、总体技术栈一览", 1)
    add_table(
        doc,
        ["层次", "技术/工具", "用途"],
        [
            ["前端", "Vue 3 + Vite", "SPA 应用与工程化构建"],
            ["前端 UI", "Element Plus / Naive UI", "管理与对话界面组件"],
            ["地图", "Leaflet（或 OpenLayers）", "二维地图渲染与交互"],
            ["后端", "Python 3.10+ / FastAPI", "API 服务、异步 IO"],
            ["AI 编排", "LangChain / LangGraph", "Agent、工具、工作流"],
            ["大模型", "国产或通用 LLM API", "意图理解与语言生成"],
            ["GIS 计算", "GeoPandas / Shapely / PyProj", "矢量分析与投影"],
            ["空间库（可选）", "PostGIS / PostgreSQL", "空间数据持久化与查询"],
            ["向量库", "Chroma / FAISS", "RAG 检索"],
            ["文档处理", "python-docx 等", "报告导出"],
            ["部署", "Docker / docker-compose（可选）", "一键演示环境"],
        ],
    )

    add_title(doc, "三、前端工具与平台", 1)
    add_title(doc, "3.1 Vue 3 + Vite", 2)
    add_para(
        doc,
        "Vue 3 提供组件化与响应式数据绑定，适合构建“对话 + 地图”双栏复杂交互。"
        "Vite 提供快速冷启动与 HMR，缩短开发迭代周期。",
    )
    add_title(doc, "3.2 地图引擎 Leaflet", 2)
    add_para(
        doc,
        "Leaflet 轻量、资料丰富，足以支撑演示所需的 GeoJSON 叠加、高亮、适应视野等能力。"
        "若后续需要更复杂专业制图，可迁移到 OpenLayers；若需要三维，可扩展 Cesium（非一期重点）。",
    )
    add_title(doc, "3.3 UI 组件库", 2)
    add_para(
        doc,
        "采用成熟 Vue 组件库快速搭建对话框、表格、抽屉、表单与消息提示，保证界面完整度与可用性。",
    )

    add_title(doc, "四、后端与服务平台", 1)
    add_title(doc, "4.1 FastAPI", 2)
    add_para(
        doc,
        "FastAPI 基于 Python 类型提示自动生成接口文档，天然适合对接 AI 与数据科学生态。"
        "支持 SSE/WebSocket，便于流式返回 Agent 中间步骤。",
    )
    add_title(doc, "4.2 任务与会话", 2)
    add_para(
        doc,
        "会话上下文可先采用内存/Redis；演示版可用本地 JSON 持久化。"
        "每次研判形成 run_id，关联工具轨迹与报告，方便复盘。",
    )

    add_title(doc, "五、人工智能相关工具平台", 1)
    add_title(doc, "5.1 LangChain", 2)
    add_para(
        doc,
        "LangChain 是面向大模型应用的编排框架，提供 Prompt 模板、输出解析、记忆、检索、工具调用与 Agent 抽象。"
        "在本项目中，LangChain 不替代 GIS，而是把 GIS 函数变成 Agent 可调用的 Tool，并管理多步推理循环。",
    )
    add_para(doc, "典型能力映射：")
    add_bullet(doc, "Tools：封装 buffer/intersect/count/retrieve 等。")
    add_bullet(doc, "Agents：根据问题自动选择工具。")
    add_bullet(doc, "Memory：支持多轮追问。")
    add_bullet(doc, "Retrievers：对接向量库完成 RAG。")

    add_title(doc, "5.2 LangGraph（进阶可选）", 2)
    add_para(
        doc,
        "当需要更强可控性时，使用 LangGraph 将流程固化为状态图："
        "槽位检查 → 空间分析 → 结果校验 → 报告生成。适合减少自由 Agent 的随机性，提高演示稳定性。",
    )

    add_title(doc, "5.3 大模型服务", 2)
    add_para(
        doc,
        "可通过 OpenAI 兼容接口接入通义、DeepSeek、文心、OpenAI 等模型。"
        "选型原则：中文理解好、工具调用稳定、成本可控、可在文档中说明降级策略。",
    )

    add_title(doc, "5.4 RAG 组件", 2)
    add_para(
        doc,
        "使用文档加载器、文本切分、Embedding、向量检索组成知识问答链路。"
        "向量库可选 Chroma（嵌入式，便于单机演示）或 FAISS。",
    )

    add_title(doc, "六、GIS 与数据工具平台", 1)
    add_table(
        doc,
        ["工具", "说明"],
        [
            ["GeoPandas", "矢量数据读写与空间连接、叠加分析的高层次接口"],
            ["Shapely", "几何对象创建与拓扑运算（缓冲、相交、包含）"],
            ["PyProj", "坐标系转换与投影"],
            ["Rasterio（可选）", "栅格数据读取，为后续遥感扩展预留"],
            ["PostGIS", "生产级空间数据库，支持空间索引与 SQL 空间函数"],
            ["QGIS（研发辅助）", "样例数据制作、检查与制图，不必然嵌入运行时"],
        ],
    )
    add_para(
        doc,
        "平台原则：运行时计算栈以 Python GIS 库为主；QGIS 作为数据准备工具；"
        "需要时可增加 GIS MCP 一类工具服务，把空间算子标准化暴露给 Agent。",
    )

    add_title(doc, "七、开发与工程化工具", 1)
    add_bullet(doc, "代码管理：Git。")
    add_bullet(doc, "Python 环境：venv / conda；依赖锁定 requirements.txt 或 poetry。")
    add_bullet(doc, "前端包管理：npm / pnpm。")
    add_bullet(doc, "接口调试：FastAPI Swagger UI、Postman/Apifox。")
    add_bullet(doc, "容器化（可选）：Docker Compose 一键启动前后端与数据库。")
    add_bullet(doc, "文档：设计 DOCX + 项目 README。")

    add_title(doc, "八、运行环境建议", 1)
    add_table(
        doc,
        ["项目", "建议配置"],
        [
            ["操作系统", "Windows 10/11 或主流 Linux"],
            ["CPU/内存", "4 核以上 / 8GB 及以上（演示）"],
            ["浏览器", "Chrome / Edge 最新稳定版"],
            ["网络", "需访问 LLM API；可准备本地模拟模式"],
            ["磁盘", "预留样例数据与模型缓存空间"],
        ],
    )

    add_title(doc, "九、平台安全与密钥管理", 1)
    add_bullet(doc, "使用 .env 管理 API Key，不提交到公开仓库。")
    add_bullet(doc, "前端仅持有业务 Token，不持有模型密钥。")
    add_bullet(doc, "上传数据大小与类型限制，防止异常文件。")
    add_bullet(doc, "演示环境默认本地回环或局域网，避免公网裸奔。")

    add_title(doc, "十、工具平台与项目目标的对应关系", 1)
    add_para(
        doc,
        "本项目聚焦大语言模型、AIGC、智能体在垂直场景落地。平台中："
        "LLM 负责理解与生成；AIGC 体现为研判报告自动生成；智能体体现为工具编排与多步空间分析；"
        "垂直场景体现为城市治理空间研判。工具平台选择均服务于“可运行、可演示、可解释”。",
    )

    add_meta_footer(doc, "05 工具平台介绍")
    path = OUT_DIR / "05_工具平台介绍.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 6. 系统架构设计
# ---------------------------------------------------------------------------
def build_architecture():
    doc = init_doc()
    cover_block(doc, "06 系统架构设计")

    add_title(doc, "一、架构目标", 1)
    add_bullet(doc, "支持自然语言驱动的空间分析闭环。")
    add_bullet(doc, "保证 GIS 计算与语言模型职责分离。")
    add_bullet(doc, "模块可替换：模型、地图、数据库可升级。")
    add_bullet(doc, "便于单机演示与后续工程扩展。")

    add_title(doc, "二、逻辑架构", 1)
    add_para(doc, "自上而下分为：")
    add_numbered(doc, "用户交互层：Web 对话、地图、报告、管理。")
    add_numbered(doc, "应用服务层：会话、鉴权、演示剧本、导出。")
    add_numbered(doc, "智能体层：LangChain/LangGraph 编排。")
    add_numbered(doc, "领域能力层：GIS Tools、RAG、Report。")
    add_numbered(doc, "数据与基础设施层：文件/PostGIS、向量库、对象存储（可选）、日志。")

    add_title(doc, "三、部署架构（演示版）", 1)
    add_para(
        doc,
        "推荐最小部署：浏览器 ↔ Vue 静态资源（或开发服务器）↔ FastAPI ↔（LLM API + 本地 GIS 数据）。"
        "可选增加 PostgreSQL/PostGIS 与 Chroma。Docker Compose 可将 api、web、db 编排为单命令启动。",
    )

    add_title(doc, "四、关键时序（问答研判）", 1)
    add_numbered(doc, "用户提交问题。")
    add_numbered(doc, "后端创建 run，进入 Agent 循环。")
    add_numbered(doc, "模型提出 tool_call；后端执行 GIS/RAG 工具。")
    add_numbered(doc, "工具结果回灌模型；必要时继续调用。")
    add_numbered(doc, "生成最终答复与 map_actions。")
    add_numbered(doc, "前端更新对话、轨迹与地图。")
    add_numbered(doc, "用户触发报告生成，写入导出文件。")

    add_title(doc, "五、数据架构", 1)
    add_table(
        doc,
        ["数据类", "存储", "说明"],
        [
            ["矢量业务数据", "GeoJSON / PostGIS", "学校医院等设施与事件点"],
            ["地名别名", "SQLite/JSON", "提高 geocode 命中"],
            ["知识文档", "文件 + 向量库", "规范与手册"],
            ["会话与轨迹", "SQLite/JSON/Redis", "演示可文件化"],
            ["配置密钥", ".env", "不入库"],
        ],
    )

    add_title(doc, "六、接口与集成原则", 1)
    add_bullet(doc, "前后端 REST/SSE；地图指令结构化。")
    add_bullet(doc, "工具输入输出 JSON Schema 化。")
    add_bullet(doc, "错误码分类：用户可修正 / 系统可重试 / 需人工介入。")
    add_bullet(doc, "观测性：run_id 贯穿日志、轨迹、报告。")

    add_title(doc, "七、扩展架构展望", 1)
    add_bullet(doc, "多 Agent：规划员、空间分析员、报告员分工。")
    add_bullet(doc, "接入政务数据中台与统一身份认证。")
    add_bullet(doc, "移动端轻应用作为外业入口。")
    add_bullet(doc, "与视频感知/物联网告警联动形成闭环工单。")

    add_meta_footer(doc, "06 系统架构设计")
    path = OUT_DIR / "06_系统架构设计.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 7. 应用案例说明
# ---------------------------------------------------------------------------
def build_cases():
    doc = init_doc()
    cover_block(doc, "07 应用案例说明")

    add_title(doc, "一、案例设计原则", 1)
    add_para(
        doc,
        "案例服务两个目的：证明垂直场景价值；支撑稳定可复现的演示。"
        "因此采用“虚构城市局部样例 + 真实业务问题类型”的方式，数据脱敏，流程真实。",
    )

    add_title(doc, "二、案例一：公共服务设施覆盖速查", 1)
    add_para(doc, "场景背景：", bold=True)
    add_para(
        doc,
        "某区城市治理专班需要快速了解“阳光花园小区”周边基础公共服务配套情况，"
        "为回应居民诉求提供数据支撑。",
    )
    add_para(doc, "用户问题示例：", bold=True)
    add_para(doc, "“阳光花园小区周边 500 米内有多少所学校和医院？请在地图上标出，并给出简要研判。”")
    add_para(doc, "系统执行要点：", bold=True)
    add_bullet(doc, "地名定位 → 500m 缓冲 → 学校/医院图层相交 → 计数与列表。")
    add_bullet(doc, "地图展示缓冲区与命中点位。")
    add_bullet(doc, "结合规范知识库给出“覆盖是否薄弱”的定性建议（谨慎表述）。")
    add_para(doc, "预期价值：", bold=True)
    add_para(doc, "将原本需要 GIS 操作的 10～20 分钟流程压缩为一次对话，结果可解释、可汇报。")

    add_title(doc, "三、案例二：投诉事件周边环境研判", 1)
    add_para(doc, "场景背景：", bold=True)
    add_para(
        doc,
        "12345/城管热线收到噪声或环境类投诉，值班人员需迅速掌握事发点周边公园、学校、居民区等要素，"
        "辅助判断影响面与优先核查方向。",
    )
    add_para(doc, "用户问题示例：", bold=True)
    add_para(doc, "“对投诉点 C2026001 做 1 公里缓冲，列出落入范围内的公园，并统计数量。”")
    add_para(doc, "系统执行要点：", bold=True)
    add_bullet(doc, "按事件编号定位坐标。")
    add_bullet(doc, "缓冲与公园图层相交，返回名称列表与数量。")
    add_bullet(doc, "生成简报，供交接班使用。")

    add_title(doc, "四、案例三：拟建点安全距离/敏感目标校核", 1)
    add_para(doc, "场景背景：", bold=True)
    add_para(
        doc,
        "在设施选址或临时点位审批中，需判断拟建点是否进入敏感目标缓冲控制范围"
        "（演示可用公园/学校等代理敏感层）。",
    )
    add_para(doc, "用户问题示例：", bold=True)
    add_para(doc, "“拟建点 P1 是否位于任一学校 300 米缓冲区内？给出判定与最近学校距离。”")
    add_para(doc, "系统执行要点：", bold=True)
    add_bullet(doc, "投影后进行米制缓冲与距离计算。")
    add_bullet(doc, "输出布尔判定、最近距离、相关学校名称。")
    add_bullet(doc, "地图高亮冲突缓冲与拟建点。")

    add_title(doc, "五、案例演示剧本", 1)
    add_table(
        doc,
        ["序号", "剧本名称", "一键问题", "成功判据"],
        [
            ["D1", "覆盖速查", "小区500m学校医院统计", "返回数量+落图+轨迹含buffer"],
            ["D2", "投诉研判", "投诉点1km公园列表", "返回列表+落图"],
            ["D3", "距离校核", "拟建点与学校300m判定", "返回是/否+距离"],
        ],
    )

    add_title(doc, "六、案例结果的业务意义", 1)
    add_para(
        doc,
        "上述案例覆盖“民生服务、事件处置、合规辅助”三类治理高频问题，"
        "体现人工智能不是替代规划制度，而是降低空间分析门槛、提升响应速度与解释能力。"
        "由此形成完整的“场景—技术—系统—案例”叙事闭环。",
    )

    add_meta_footer(doc, "07 应用案例说明")
    path = OUT_DIR / "07_应用案例说明.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 8. 演示与部署说明
# ---------------------------------------------------------------------------
def build_deploy():
    doc = init_doc()
    cover_block(doc, "08 演示与部署说明")

    add_title(doc, "一、演示目标", 1)
    add_para(
        doc,
        "在 5～10 分钟内完整展示：自然语言提问、Agent 调用 GIS 工具、地图结果联动、"
        "工具轨迹可解释、研判报告生成。",
    )

    add_title(doc, "二、环境准备清单", 1)
    add_bullet(doc, "安装 Python 3.10+、Node.js 18+。")
    add_bullet(doc, "配置 LLM API Key（.env）。")
    add_bullet(doc, "准备样例 GeoJSON 与演示剧本配置。")
    add_bullet(doc, "浏览器 Chrome/Edge。")

    add_title(doc, "三、启动步骤（规划）", 1)
    add_numbered(doc, "启动后端：安装依赖后运行 uvicorn/fastapi 服务。")
    add_numbered(doc, "启动前端：npm install && npm run dev。")
    add_numbered(doc, "浏览器打开本地地址，点击 Demo 剧本。")
    add_numbered(doc, "展示自由问答与报告导出。")
    add_para(doc, "（具体命令以项目 README 最终实现为准。）")

    add_title(doc, "四、演示话术建议", 1)
    add_bullet(doc, "先讲痛点：业务人员不会复杂 GIS。")
    add_bullet(doc, "再讲范式：LLM 编排 + GIS 工具保真。")
    add_bullet(doc, "先点 D1/D2/D3，强调轨迹面板。")
    add_bullet(doc, "最后展示报告，对应 AIGC 产出。")

    add_title(doc, "五、风险与预案", 1)
    add_table(
        doc,
        ["风险", "预案"],
        [
            ["模型 API 波动", "预设缓存答复 / 本地模拟 Agent"],
            ["网络受限", "局域网部署 + 事先下载依赖"],
            ["地名匹配失败", "使用剧本固定名称与地图选点"],
            ["操作失误", "一键 Demo，避免临时手输"],
        ],
    )

    add_title(doc, "六、交付检查表", 1)
    add_bullet(doc, "设计文档分册齐全。")
    add_bullet(doc, "系统可运行。")
    add_bullet(doc, "案例可演示。")
    add_bullet(doc, "技术说明与启动方式清晰。")

    add_meta_footer(doc, "08 演示与部署说明")
    path = OUT_DIR / "08_演示与部署说明.docx"
    doc.save(path)
    return path


# ---------------------------------------------------------------------------
# 9. 文档目录说明
# ---------------------------------------------------------------------------
def build_index():
    doc = init_doc()
    cover_block(doc, "00 设计文档目录与说明")

    add_title(doc, "一、文档体系说明", 1)
    add_para(
        doc,
        "为完整说明项目设计，本仓库将设计文档拆分为独立 DOCX 分册，便于查阅与版本管理。"
        "各分册互补，共同构成完整设计说明，覆盖需求分析、设计思路、功能模块和工具平台介绍等内容。",
    )

    add_title(doc, "二、分册清单", 1)
    add_table(
        doc,
        ["文件名", "内容"],
        [
            ["00_设计文档目录与说明.docx", "文档索引与阅读指引"],
            ["01_项目概述.docx", "背景、定位、目标、创新点、交付物"],
            ["02_需求分析.docx", "角色、功能/非功能需求、用户故事、验收"],
            ["03_设计思路.docx", "理念、抽象、技术路径、权衡与演示设计"],
            ["04_功能模块.docx", "模块划分、协作流程、接口与验收"],
            ["05_工具平台介绍.docx", "前后端、AI、GIS、工程化工具说明"],
            ["06_系统架构设计.docx", "逻辑/部署/数据架构与时序"],
            ["07_应用案例说明.docx", "三类治理案例与演示剧本"],
            ["08_演示与部署说明.docx", "演示流程、环境与预案"],
        ],
    )

    add_title(doc, "三、建议阅读顺序", 1)
    add_para(doc, "概述 → 需求 → 设计思路 → 功能模块 → 工具平台 → 架构 → 案例 → 演示。")

    add_title(doc, "四、与软件系统的关系", 1)
    add_para(
        doc,
        "设计文档描述“做什么、为何如此做、模块与平台是什么”；"
        "可运行 Web 系统证明“已经做成并可演示”。二者需保持模块命名与案例剧本一致。",
    )

    add_meta_footer(doc, "00 设计文档目录与说明")
    path = OUT_DIR / "00_设计文档目录与说明.docx"
    doc.save(path)
    return path


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = [
        build_index(),
        build_overview(),
        build_requirements(),
        build_design(),
        build_modules(),
        build_platform(),
        build_architecture(),
        build_cases(),
        build_deploy(),
    ]
    print("OUT_DIR=", OUT_DIR)
    for p in paths:
        print("OK", p.name, p.stat().st_size)


if __name__ == "__main__":
    main()
